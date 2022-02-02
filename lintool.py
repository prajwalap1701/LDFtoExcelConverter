import os
from tkinter import *
from tkinter import filedialog, messagebox

import ldfparser
from openpyxl import Workbook
from openpyxl.worksheet.table import Table
import re
import math

filepath = ''
output_file_name = ''

def setTextInput(text):
    global e1
    e1.delete(0,"end")
    e1.insert(0, text)
    return

def browseFiles():
    try:
        filename = filedialog.askopenfilename(initialdir="/", title="Select a LDF File",
                                              filetypes=(("LDF files", "*.ldf*"), ("all files", "*.*")))
        global filepath, output_file_name
        filepath = filename
        output_file_name = re.split("/", filepath)[-1].replace('.ldf', '.xlsx')
        setTextInput(str(filepath))


    except:
        messagebox.showerror("Error", "Unable to load File")

def convert_ldf():
    try:
        print('In convert_ldf')
        wb = Workbook()

        # input file path
        ldf = ldfparser.parse_ldf_to_dict(filepath)

        print('Stage 0')

        frames = ldf['frames']
        signals = ldf['signals']

        fr_sig_list = []

        for f in frames:
            for s in f['signals']:
                for sig in signals:
                    if sig['name'] == s['signal']:
                        s['width'] = sig['width']
                        s['init_value'] = sig['init_value']
                        s['publisher'] = sig['publisher']
                        s['subscribers'] = sig['subscribers']
            fr_sig_list.append(f)

        data1 = []
        for f in fr_sig_list:
            for s in f['signals']:
                data1.append(['0x' + (format(int(f['frame_id']), 'X')), s['signal'], '', f['name'], f['publisher'],
                              re.sub('\[|\]|\'', '', str(s['subscribers'])), 'Intel', s['width'],
                              s['offset'], '', '', '', '', s['init_value'], '', '', '', '', '', '-'])
        print('Stage 1')

        # *********************************************
        # Create new list in this format [encoding_name, encoding value] and add to data1 list
        temp_str = ''
        encoding_list = []
        phy_encoding_list = []

        for enc in ldf['signal_encoding_types']:
            for val in enc['values']:
                if val['type'] == 'physical':
                    phy_encoding_list.append([enc['name'], val])
                if val['type'] == 'logical':
                    temp_str += str(val['value']) + " \"" + str(val['text']) + "\" "
            encoding_list.append([enc['name'], temp_str])
            temp_str = ''

        for d in data1:
            for e in encoding_list:
                if d[1] in e[0]:
                    d[18] = e[1]

        # Fills offset, min, max etc
        for d in data1:
            for p in phy_encoding_list:
                if d[1] in p[0]:
                    d[11] = p[1]['scale']
                    d[12] = p[1]['offset']
                    d[14] = p[1]['min']
                    d[15] = p[1]['max']
                    d[17] = p[1]['unit']
                else:
                    d[11] = 1
                    d[12] = 0
                    d[14] = 0
                    d[15] = math.pow(2, d[7]) - 1

        print('Stage 2')

        # **********************************************

        ws = wb.create_sheet('LIN Signal')
        sheet = wb["Sheet"]
        wb.remove(sheet)

        ws.append(['Identifier', 'Signal Name', 'Signal Description', 'Frames',
                   'Publisher', 'Subscribers', 'Byte Order', 'Data length [bits]',
                   'Start bit', 'DataType', 'Signal type',
                   'Factor', 'Offset', 'Initial value [dez]', 'Minimum (phys)', 'Maximum (phys)',
                   'Invalid value [hex]', 'Unit', 'Logical Encoding', 'Comment'])

        for row in data1:
            ws.append(row)
        ref_str = "A1" + ":" + "T" + str(len(data1) + 1)
        tab1 = Table(displayName="Table1", ref=ref_str)

        print('Stage 3')


        ws.add_table(tab1)

        wb.save(os.path.expanduser('~\\Documents\\' + output_file_name))
        messagebox.showinfo("Export Successful",
                            "File saved to " + os.path.expanduser('~\\Documents\\' + output_file_name))

        os._exit(0)

    except:
        messagebox.showerror("Error", "Unable to Export!! Inconsistent LDF file")
        os._exit(0)


master = Tk()
master.title("LDF to Excel Converter")
master.minsize(width=400, height=50)

l1=Label(master, text="Input LDF File").grid(row=0)

e1 = Entry(master, width=50)

e1.grid(row=0, column=1)

button_explore = Button(master, text="Browse File", command= browseFiles).grid(row=0, column=2)

Button(master, text='Convert',command=convert_ldf).grid(row=4, column=1,pady=4)

mainloop()