import ldfparser
from openpyxl import Workbook
from openpyxl.worksheet.table import Table
import re
import math

wb = Workbook()

# input file path
ldf = ldfparser.parse_ldf_to_dict('bcm.ldf')

frames = ldf['frames']
signals = ldf['signals']

fr_sig_list = []

for f in frames:
    for s in f['signals']:
        for sig in signals:
            if sig['name'] == s['signal']:
                s['width'] = sig['width']
                s['init_value'] = sig['init_value']
                s['publisher'] = sig['publisher']
                s['subscribers'] = sig['subscribers']
    fr_sig_list.append(f)


data1 = []
for f in fr_sig_list:
    for s in f['signals']:
        data1.append(['0x'+(format(int(f['frame_id']), 'X')), s['signal'], '', f['name'], f['publisher'],
                      re.sub('\[|\]|\'', '', str(s['subscribers'])), 'Intel', s['width'],
                      s['offset'], '', '', '', '', s['init_value'], '', '', '', '', '', '-'])

# *********************************************
# Create new list in this format [encoding_name, encoding value] and add to data1 list
temp_str = ''
encoding_list = []
phy_encoding_list = []

for enc in ldf['signal_encoding_types']:
    for val in enc['values']:
        if val['type'] == 'physical':
            phy_encoding_list.append([enc['name'], val])
        if val['type'] == 'logical':
            temp_str += str(val['value']) + " \""+str(val['text'])+ "\" "
    encoding_list.append([enc['name'], temp_str])
    temp_str = ''

for d in data1:
    for e in encoding_list:
        if d[1] in e[0]:
            d[18] = e[1]

# Fills offset, min, max etc
for d in data1:
    for p in phy_encoding_list:
        if d[1] in p[0]:
            d[11] = p[1]['scale']
            d[12] = p[1]['offset']
            d[14] = p[1]['min']
            d[15] = p[1]['max']
            d[17] = p[1]['unit']
        else:
            d[11] = 1
            d[12] = 0
            d[14] = 0
            d[15] = math.pow(2, d[7]) - 1


# **********************************************

ws = wb.create_sheet('LIN Signal')
sheet = wb["Sheet"]
wb.remove(sheet)

ws.append(['Identifier', 'Signal Name', 'Signal Description', 'Frames',
           'Publisher', 'Subscribers', 'Byte Order', 'Data length [bits]',
           'Start bit', 'DataType', 'Signal type',
           'Factor', 'Offset', 'Initial value [dez]', 'Minimum (phys)', 'Maximum (phys)',
           'Invalid value [hex]', 'Unit', 'Logical Encoding', 'Comment'])

for row in data1:
    ws.append(row)
ref_str = "A1" + ":" + "T" + str(len(data1) + 1)
tab1 = Table(displayName="Table1", ref=ref_str)

ws.add_table(tab1)

# output file path
wb.save('new.xlsx')